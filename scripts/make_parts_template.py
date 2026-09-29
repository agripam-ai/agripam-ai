#!/usr/bin/env python3
"""Create the blank BioBrick parts workbook (data/biobricks/Parts_input_template.xlsx)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "data/biobricks/Parts_input_template.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "Parts"
header = ["part_id", "name", "role", "sequence", "hosts", "source", "notes"]
ws.append(header)
for cell in ws[1]:
    cell.fill, cell.font = PatternFill("solid", fgColor="2E9E6B"), Font(bold=True, color="FFFFFF")
ws.append(["EXAMPLE_delete_this_row", "example promoter", "promoter", "TTGACGGCTAGCTCAGTCCTAGGTACAGTGCTAGC",
           "E. coli; B. subtilis", "where the sequence comes from", "rows starting with EXAMPLE are ignored"])
for col, width in zip("ABCDEFG", (26, 30, 14, 60, 26, 40, 40)):
    ws.column_dimensions[col].width = width
dv = DataValidation(type="list", formula1='"promoter,rbs,cds,reporter,terminator,scaffold,backbone,other"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("C2:C500")
info = wb.create_sheet("Instructions")
for line in ("One row per part.",
             "part_id: a BioBrick ID (BBa_...) or your own code. name: free text.",
             "role: promoter, rbs, cds, reporter, terminator, scaffold (sgRNA), backbone or other.",
             "sequence: A, C, G, T only, in the 5' to 3' direction, without the BioBrick prefix and suffix.",
             "hosts: species where the part is known to work, separated by semicolons (for example E. coli; B. subtilis). "
             "Leave blank if unknown: the app then warns that activity in your chassis is untested.",
             "You can also fetch parts by BBa_ ID from the iGEM Registry inside the app (needs internet)."):
    info.append([line])
info.column_dimensions["A"].width = 130
for row in info.iter_rows():
    row[0].alignment = Alignment(wrap_text=True)
wb.save(OUT)
print(OUT)
