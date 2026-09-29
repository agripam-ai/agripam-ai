#!/usr/bin/env python3
"""Convert a simple Markdown file (headings, paragraphs, bullets, numbered lists, pipe tables, **bold**, [link](url)) to .docx.

Usage: python scripts/md_to_docx.py input.md output.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_final_deliverables import new_doc, shade, table  # noqa: E402

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|\[[^\]]+\]\([^)]+\))")


def add_runs(paragraph, text: str, size: float | None = None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            run = paragraph.add_run(f"{label} ({url})")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1]); run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def convert(src: Path, dst: Path):
    lines = src.read_text().splitlines()
    title = next((l[2:] for l in lines if l.startswith("# ")), src.stem)
    subtitle = next((l for l in lines[1:] if l.strip() and not l.startswith("#")), "")
    d = new_doc(title, subtitle)
    i, skipped_sub = 0, False
    while i < len(lines):
        line = lines[i]
        if line.startswith("# ") and line[2:] == title and not skipped_sub:
            skipped_sub = True
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            i += 1  # the subtitle line
            continue
        if line.startswith("# "):
            d.add_heading(line[2:], 1)
        elif line.startswith("## "):
            d.add_heading(line[3:], 2)
        elif line.startswith("### "):
            d.add_heading(line[4:], 3)
        elif re.match(r"!\[.*\]\(.+\)\s*$", line):
            cap, path = re.match(r"!\[(.*)\]\((.+)\)\s*$", line).groups()
            img = (src.parent / path)
            if img.exists():
                d.add_picture(str(img), width=Inches(6.3))
            cp = d.add_paragraph(); add_runs(cp, cap, 9)
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            header, rows = block[0], [r for r in block[2:]]
            t = table(d, header, [[re.sub(r"\*\*|\[([^\]]+)\]\(([^)]+)\)", lambda m: (m.group(1) + " (" + m.group(2) + ")") if m.group(1) else "", c) for c in r] for r in rows], None, 8.5)
            continue
        elif re.match(r"\s*[-*] ", line):
            p = d.add_paragraph(style="List Bullet"); add_runs(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"\d+\. ", line):
            p = d.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\d+\. ", "", line))
        elif line.strip():
            p = d.add_paragraph(); add_runs(p, line)
        i += 1
    d.save(dst)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
